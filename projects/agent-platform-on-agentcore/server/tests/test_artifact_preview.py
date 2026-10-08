"""A light preview for binary artifacts.

Extraction is best-effort by design: an unparseable .docx, a missing library or a
format with no extractor must all degrade to a plain download card, never to a
failed artifact. The artifact is already stored and downloadable by the time this
runs.
"""
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from services.artifact_preview_service import preview_for  # noqa: E402


def docx_bytes(paragraphs):
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes(rows):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "결과"
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def rich_docx_bytes():
    """A document with every structure the extractor is expected to carry over."""
    from docx import Document

    document = Document()
    document.add_heading("보고서 제목", level=1)
    document.add_heading("1. 개요", level=2)
    paragraph = document.add_paragraph("이 문서는 ")
    paragraph.add_run("굵은 글씨").bold = True
    paragraph.add_run(" 와 ")
    paragraph.add_run("기울임").italic = True
    paragraph.add_run(" 을 포함한다.")
    document.add_paragraph("첫 번째 항목", style="List Bullet")
    document.add_paragraph("두 번째 항목", style="List Bullet")
    document.add_paragraph("하나", style="List Number")
    table = document.add_table(rows=2, cols=3)
    for row, cells in enumerate([["구분", "금액", "비고"], ["A", "100", "-"]]):
        for column, value in enumerate(cells):
            table.cell(row, column).text = value
    document.add_paragraph("표 뒤 문단")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes():
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "분기 실적"
    body = slide.placeholders[1].text_frame
    body.text = "매출 성장"
    nested = body.add_paragraph()
    nested.text = "북미 중심"
    nested.level = 1
    slide.notes_slide.notes_text_frame.text = "발표자 메모"

    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "표"
    shape = second.shapes.add_table(
        2, 2, Inches(1), Inches(2), Inches(4), Inches(1)
    )
    for row, cells in enumerate([["항목", "값"], ["A", "1"]]):
        for column, value in enumerate(cells):
            shape.table.cell(row, column).text = value

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_docx_preview_is_markdown_with_the_paragraphs():
    text, media = preview_for("보고서.docx", docx_bytes(["제목", "본문 첫 문단"]))
    assert media == "markdown"
    assert "제목" in text
    assert "본문 첫 문단" in text


def test_docx_headings_carry_their_level():
    text, _ = preview_for("보고서.docx", rich_docx_bytes())
    assert "# 보고서 제목" in text
    assert "## 1. 개요" in text


def test_docx_runs_keep_bold_and_italic():
    """Emphasis has to end up as markdown syntax, not be flattened into plain text."""
    text, _ = preview_for("보고서.docx", rich_docx_bytes())
    assert "**굵은 글씨**" in text
    assert "*기울임*" in text
    # The markers must hug the words: `** 굵은 글씨 **` is literal asterisks.
    assert "** 굵은" not in text


def test_docx_list_paragraphs_become_markdown_lists():
    text, _ = preview_for("보고서.docx", rich_docx_bytes())
    assert "- 첫 번째 항목" in text
    assert "- 두 번째 항목" in text
    assert "1. 하나" in text


def test_docx_table_has_the_separator_row_gfm_needs():
    """Pipe-joined lines without a separator row render as paragraphs of text."""
    text, _ = preview_for("보고서.docx", rich_docx_bytes())
    assert "| 구분 | 금액 | 비고 |" in text
    assert "| --- | --- | --- |" in text
    assert "| A | 100 | - |" in text


def test_docx_keeps_document_order():
    """A table is a body element like any other, so it cannot be appended at the end."""
    text, _ = preview_for("보고서.docx", rich_docx_bytes())
    assert text.index("| 구분") < text.index("표 뒤 문단")


def test_pptx_preview_is_markdown_per_slide():
    text, media = preview_for("실적.pptx", pptx_bytes())
    assert media == "markdown"
    assert "## 1. 분기 실적" in text
    assert "- 매출 성장" in text
    # Outline level becomes list indentation.
    assert "  - 북미 중심" in text
    assert "> 발표자 메모" in text
    # The title is the slide heading, not also a bullet under it.
    assert "- 분기 실적" not in text
    assert "## 2. 표" in text
    assert "| 항목 | 값 |" in text
    assert "| --- | --- |" in text


def test_corrupt_pptx_does_not_raise():
    assert preview_for("실적.pptx", b"not a zip at all") is None


def test_xlsx_preview_is_sheets_with_sheets_and_rows():
    import json

    text, media = preview_for("결과.xlsx", xlsx_bytes([["이름", "값"], ["가", 1]]))
    assert media == "sheets"
    sheets = json.loads(text)["sheets"]
    assert sheets[0]["name"] == "결과"
    assert sheets[0]["rows"][0] == ["이름", "값"]
    assert sheets[0]["rows"][1] == ["가", "1"]
    assert sheets[0]["truncated"] is False


def test_csv_and_markdown_files_preview_as_their_own_text():
    text, media = preview_for("data.csv", "이름,값\n가,1\n".encode("utf-8"))
    assert media == "csv"
    assert text.startswith("이름,값")


def test_markdown_files_have_media_name_markdown():
    text, media = preview_for("README.md", "# 제목\n\n본문\n".encode("utf-8"))
    assert media == "markdown"
    assert "제목" in text


def test_text_files_have_media_name_text():
    text, media = preview_for("notes.txt", "메모\n더 많은 메모\n".encode("utf-8"))
    assert media == "text"
    assert "메모" in text


def test_a_format_with_no_extractor_has_no_preview():
    assert preview_for("archive.zip", b"PK\x03\x04") is None
    assert preview_for("report.pdf", b"%PDF-1.7") is None   # rendered by the browser


def test_corrupt_input_does_not_raise():
    assert preview_for("보고서.docx", b"not a zip at all") is None


def test_preview_is_truncated():
    text, _ = preview_for("big.csv", ("x" * 200_000).encode("utf-8"))
    assert len(text) <= 100_000


def test_large_workbook_preview_stays_valid_json():
    """A workbook bigger than PREVIEW_MAX_CHARS is sliced at the data level, not the JSON.

    Before the fix, slicing the serialised JSON string at 100,000 chars cuts it mid-token,
    leaving the preview unparseable.
    """
    import json

    # Build a 200×20 workbook with repeated strings to exceed PREVIEW_MAX_CHARS
    rows = [
        [f"셀값{r}-{c}" * 12 for c in range(20)]
        for r in range(200)
    ]
    body = xlsx_bytes(rows)
    text, media = preview_for("big.xlsx", body)

    assert media == "sheets"
    assert len(text) <= 100_000, f"Preview too large: {len(text)} chars"

    # The critical check: the JSON must be valid
    sheets = json.loads(text)["sheets"]
    assert len(sheets) > 0, "Preview should have at least one sheet"
    assert len(sheets[0]["rows"]) > 0, "Sheet should have been truncated, not empty"
    # The panel says so on screen — a silently shortened sheet reads as the
    # whole workbook.
    assert sheets[0]["truncated"] is True
