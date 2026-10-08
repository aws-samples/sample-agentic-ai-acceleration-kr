"""Extracts a light, readable preview from a binary artifact.

Everything here is best-effort. The artifact is already stored and downloadable
before extraction runs, so an unparseable document, a missing library or a format
with no extractor must all leave the artifact intact and simply produce no
preview — the panel then shows a plain download card, which is also the steady
state for .zip and .pdf.

PDF has no extractor on purpose: the panel renders it in the browser from the
file's own bytes, and a text dump would be worse than the real thing. docx and
pptx are rendered the same way, so their extractors here are only the fallback for
when that render fails.
"""
import json
import logging
import re
import threading
from typing import Any, Iterable, List, Optional, Tuple

from models.artifact import ArtifactVersion

logger = logging.getLogger(__name__)

PREVIEW_MAX_CHARS = 100_000
XLSX_MAX_ROWS = 1_000

_HEADING_STYLE = re.compile(r"heading\s*(\d)", re.IGNORECASE)


def _gfm_table(rows: List[List[str]]) -> str:
    """Rows as a GFM table, separator row included.

    Without the separator every flavour of markdown renders the pipes as
    literal text, which is what made docx tables read as garbage.
    """
    rows = [row for row in rows if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [list(row) + [""] * (width - len(row)) for row in rows]
    lines = [
        "| " + " | ".join(padded[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines += ["| " + " | ".join(row) + " |" for row in padded[1:]]
    return "\n".join(lines)


def _cell_markdown(text: str) -> str:
    # A pipe would end the column early and a newline the whole row.
    return text.strip().replace("|", "\\|").replace("\n", "<br>")


def _emphasised(run: Any) -> str:
    """One run with its bold/italic carried over as markdown markers."""
    text = run.text
    if not text.strip():
        return text
    marker = ("**" if run.bold else "") + ("*" if run.italic else "")
    if not marker:
        return text
    # The markers cannot wrap the run's outer whitespace — `** word **` is
    # literal asterisks in every flavour — so the padding stays outside.
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()) :]
    return f"{lead}{marker}{text.strip()}{marker}{trail}"


def _list_prefix(paragraph: Any, style: str) -> Optional[str]:
    """`- ` / `1. ` with indentation, or None when this is not a list item.

    Numbering lives in `w:numPr`, not in the style name: a list made through
    Word's UI is styled `List Paragraph`, so keying off the style alone misses
    most real documents.
    """
    numbering = None
    properties = paragraph._p.pPr
    if properties is not None:
        numbering = properties.numPr
    if numbering is None and not style.lower().startswith("list"):
        return None
    level = 0
    if numbering is not None and numbering.ilvl is not None:
        level = numbering.ilvl.val or 0
    else:
        trailing = style.rsplit(" ", 1)[-1]
        if trailing.isdigit():
            level = int(trailing) - 1
    marker = "1." if "number" in style.lower() else "-"
    return "  " * max(level, 0) + marker + " "


def _paragraph_markdown(paragraph: Any) -> str:
    text = "".join(_emphasised(run) for run in paragraph.runs).strip()
    if not text:
        return ""
    style = (getattr(paragraph.style, "name", None) or "").strip()
    if style in ("Title", "Subtitle"):
        return f"{'#' if style == 'Title' else '##'} {paragraph.text.strip()}"
    heading = _HEADING_STYLE.search(style)
    if heading:
        # Emphasis markers inside a heading are noise, so the plain text wins.
        return f"{'#' * int(heading.group(1))} {paragraph.text.strip()}"
    prefix = _list_prefix(paragraph, style)
    return f"{prefix}{text}" if prefix else text


def _docx_blocks(document: Any) -> Iterable[Any]:
    """Paragraphs and tables in document order.

    `document.paragraphs` and `document.tables` are two separate flat lists, so
    reading them in turn moves every table to the end of the preview.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def _docx_preview(body: bytes) -> str:
    import io

    from docx import Document
    from docx.table import Table

    document = Document(io.BytesIO(body))
    blocks = []
    for block in _docx_blocks(document):
        if isinstance(block, Table):
            table = _gfm_table(
                [[_cell_markdown(cell.text) for cell in row.cells] for row in block.rows]
            )
            if table:
                blocks.append(table)
            continue
        rendered = _paragraph_markdown(block)
        if rendered:
            blocks.append(rendered)
    # Blank lines between list items only make the list loose, which still
    # renders as a list — worth it to keep the joining rule this simple.
    return "\n\n".join(blocks)


def _pptx_shapes(shapes: Any) -> Iterable[Any]:
    """Shapes with groups flattened — a grouped text box holds text too."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except Exception:
            # shape_type raises for a few placeholder kinds rather than
            # answering; those are never groups.
            is_group = False
        if is_group:
            yield from _pptx_shapes(shape.shapes)
        else:
            yield shape


def _pptx_preview(body: bytes) -> str:
    import io

    from pptx import Presentation

    presentation = Presentation(io.BytesIO(body))
    blocks = []
    for index, slide in enumerate(presentation.slides, start=1):
        try:
            title_shape = slide.shapes.title
        except Exception:
            title_shape = None
        title = (title_shape.text.strip() if title_shape is not None else "") or None
        blocks.append(f"## {index}. {title}" if title else f"## Slide {index}")
        # By id, not identity: iterating the collection builds a fresh proxy per
        # shape, so the title placeholder is never the same object twice.
        title_id = title_shape.shape_id if title_shape is not None else None
        for shape in _pptx_shapes(slide.shapes):
            if title_id is not None and shape.shape_id == title_id:
                continue
            if getattr(shape, "has_table", False):
                table = _gfm_table(
                    [
                        [_cell_markdown(cell.text) for cell in row.cells]
                        for row in shape.table.rows
                    ]
                )
                if table:
                    blocks.append(table)
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            lines = []
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    lines.append("  " * max(paragraph.level or 0, 0) + f"- {text}")
            if lines:
                blocks.append("\n".join(lines))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(
                    "\n".join(f"> {line}" for line in notes.splitlines() if line.strip())
                )
    return "\n\n".join(blocks)


def _xlsx_preview(body: bytes) -> str:
    import io

    from openpyxl import load_workbook

    # read_only keeps a large workbook from being materialised in full.
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = []
        capped = False
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= XLSX_MAX_ROWS:
                capped = True
                break
            rows.append(["" if value is None else str(value) for value in row])
        sheets.append({"name": sheet.title, "rows": rows, "truncated": capped})
    text = json.dumps({"sheets": sheets}, ensure_ascii=False)
    while len(text) > PREVIEW_MAX_CHARS and any(sheet["rows"] for sheet in sheets):
        # Halve every sheet's rows and re-serialise until it fits. Slicing the
        # JSON string instead would hand the frontend a document that ends
        # mid-token — the panel parses this, so it has to stay parseable.
        for sheet in sheets:
            sheet["rows"] = sheet["rows"][: len(sheet["rows"]) // 2]
            sheet["truncated"] = True
        text = json.dumps({"sheets": sheets}, ensure_ascii=False)
    return text


def preview_for(filename: str, body: bytes) -> Optional[Tuple[str, str]]:
    """`(text, media)` for a previewable file, or None.

    `media` names how the frontend should render the text. The valid names are:
    - `markdown`: for docx, pptx and md files
    - `sheets`: for xlsx files
    - `csv`: for csv files
    - `text`: for txt files
    - `html`: for html files
    - `json`: for json files
    """
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if extension == "docx":
            return _docx_preview(body)[:PREVIEW_MAX_CHARS], "markdown"
        if extension == "pptx":
            return _pptx_preview(body)[:PREVIEW_MAX_CHARS], "markdown"
        if extension == "xlsx":
            return _xlsx_preview(body), "sheets"
        if extension in ("csv", "md", "txt", "html", "json"):
            text = body.decode("utf-8", "replace")[:PREVIEW_MAX_CHARS]
            # The frontend switches on these names, so `md` and `txt` cannot pass
            # through verbatim — they must map to the standardized names.
            media = {"md": "markdown", "txt": "text"}.get(extension, extension)
            return text, media
    except Exception as exc:
        logger.info("No preview for %s: %s", filename, exc)
        return None
    return None


class ArtifactPreviewService:
    """Writes a preview beside its artifact and records the key on the row."""

    def __init__(self, artifact_service: Any):
        self.artifact_service = artifact_service

    def extract(self, record: ArtifactVersion) -> Optional[str]:
        service = self.artifact_service
        if not getattr(service, "enabled", False) or not record.filename:
            return None
        try:
            body = service.s3.get_object(
                Bucket=service.bucket, Key=record.s3_key
            )["Body"].read()
            extracted = preview_for(record.filename, body)
            if not extracted:
                return None
            text, media = extracted
            # A sibling of the object it describes, so lifecycle rules and any
            # future cleanup treat the two together.
            preview_key = f"{record.s3_key.rsplit('.', 1)[0]}.preview.{media}"
            service.s3.put_object(
                Bucket=service.bucket,
                Key=preview_key,
                Body=text.encode("utf-8"),
                ContentType="text/plain; charset=utf-8",
            )
            service.repository.set_preview_key(
                record.artifact_id, record.version, preview_key
            )
            return preview_key
        except Exception:
            logger.warning(
                "Preview extraction failed for %s v%s",
                record.artifact_id,
                record.version,
                exc_info=True,
            )
            return None

    def spawn(self, record: ArtifactVersion) -> None:
        """Detached: extraction happens after the artifact is already deliverable.

        Never on the stream's critical path — there are no heartbeats after the
        turn's loop, and a large workbook can take seconds.
        """
        threading.Thread(
            target=self.extract,
            args=(record,),
            name=f"artifact-preview-{record.artifact_id}-v{record.version}",
            daemon=True,
        ).start()
